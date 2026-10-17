"""Read and write text in supported document formats."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Callable
from xml.etree import ElementTree as ET

SUPPORTED_EXTENSIONS = {".txt", ".csv", ".docx", ".xlsx", ".pdf", ".flextext", ".xml"}

Transform = Callable[[str], str]


def detect_ext(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS and ext != ".flextext":
        raise ValueError(
            f"Неподдерживаемый формат: {ext or '(без расширения)'}. "
            f"Допустимо: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return ext


def process_file(filename: str, data: bytes, transform: Transform) -> tuple[bytes, str, str]:
    """
    Apply transform to textual content; return (out_bytes, out_filename, media_type).
    """
    ext = detect_ext(filename)
    stem = Path(filename).stem
    out_name = f"{stem}_converted{ext}"

    if ext == ".txt":
        text = data.decode("utf-8", errors="replace")
        out = transform(text).encode("utf-8")
        return out, out_name, "text/plain; charset=utf-8"

    if ext == ".csv":
        return _process_csv(data, transform, out_name)

    if ext == ".docx":
        return _process_docx(data, transform, out_name)

    if ext == ".xlsx":
        return _process_xlsx(data, transform, out_name)

    if ext == ".pdf":
        return _process_pdf(data, transform, f"{stem}_converted.pdf")

    if ext in (".flextext", ".xml"):
        return _process_flextext(data, transform, out_name)

    raise ValueError(f"Неподдерживаемый формат: {ext}")


def _process_csv(data: bytes, transform: Transform, out_name: str) -> tuple[bytes, str, str]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    for row in reader:
        writer.writerow([transform(cell) for cell in row])
    return buf.getvalue().encode("utf-8"), out_name, "text/csv; charset=utf-8"


def _process_docx(data: bytes, transform: Transform, out_name: str) -> tuple[bytes, str, str]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    for p in doc.paragraphs:
        if p.text:
            _replace_paragraph_text(p, transform(p.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text:
                        _replace_paragraph_text(p, transform(p.text))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), out_name, (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    """Replace paragraph text while trying to keep the first run's style."""
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def _process_xlsx(data: bytes, transform: Transform, out_name: str) -> tuple[bytes, str, str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str):
                    cell.value = transform(cell.value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), out_name, (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _process_pdf(data: bytes, transform: Transform, out_name: str) -> tuple[bytes, str, str]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    raw = "\n".join(parts)
    converted = transform(raw)

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError:
        # reportlab not installed — return plain text instead of PDF
        txt_name = out_name.rsplit(".", 1)[0] + ".txt"
        return converted.encode("utf-8"), txt_name, "text/plain; charset=utf-8"

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    font_name = _register_unicode_font()
    c.setFont(font_name, 11)
    y = height - 40
    left = 40
    max_width = width - 80
    for line in converted.splitlines() or [""]:
        while line:
            fit = line
            while fit and c.stringWidth(fit, font_name, 11) > max_width:
                fit = fit[:-1]
            if not fit:
                fit = line[:1]
            c.drawString(left, y, fit)
            line = line[len(fit) :].lstrip() if fit != line else ""
            y -= 14
            if y < 40:
                c.showPage()
                c.setFont(font_name, 11)
                y = height - 40
        if not converted.splitlines():
            break
    c.save()
    return buf.getvalue(), out_name, "application/pdf"


_FONT_REGISTERED: str | None = None


def _register_unicode_font() -> str:
    global _FONT_REGISTERED
    if _FONT_REGISTERED:
        return _FONT_REGISTERED
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("UrmiSans", path))
                _FONT_REGISTERED = "UrmiSans"
                return _FONT_REGISTERED
            except Exception:
                continue
    _FONT_REGISTERED = "Helvetica"
    return _FONT_REGISTERED


def _process_flextext(data: bytes, transform: Transform, out_name: str) -> tuple[bytes, str, str]:
    """
    FLEx .flextext / .xml: transform only transcription items (<item type="txt">).
    Glosses, notes and titles are left unchanged.
    """
    raw = data.decode("utf-8", errors="replace")
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        out = transform(raw).encode("utf-8")
        return out, Path(out_name).with_suffix(".txt").name, "text/plain; charset=utf-8"

    for elem in root.iter():
        if _local(elem.tag) != "item":
            continue
        if (elem.get("type") or "").lower() != "txt":
            continue
        _transform_nested_text(elem, transform)

    xml_body = ET.tostring(root, encoding="unicode")
    if not xml_body.startswith("<?xml"):
        xml_body = '<?xml version="1.0" encoding="UTF-8"?>\n' + xml_body
    return xml_body.encode("utf-8"), out_name, "application/xml; charset=utf-8"


def _transform_nested_text(elem, transform: Transform) -> None:
    """Convert this element's text and every nested text node."""
    if elem.text:
        elem.text = transform(elem.text)
    for child in list(elem):
        _transform_nested_text(child, transform)
        if child.tail:
            child.tail = transform(child.tail)


def _local(tag: str) -> str:
    if "}" in tag:
        return tag.rsplit("}", 1)[-1]
    return tag
